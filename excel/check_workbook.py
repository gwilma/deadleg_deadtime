"""Recalculate excel/waterhammer.xlsx with LibreOffice for several cases and
compare it with the Python model (MOC run with the workbook's 10 reaches).

    python excel/check_workbook.py      (needs openpyxl and LibreOffice `soffice`)
"""

from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from waterhammer import PRESETS, VALVE_CLOSURE_TIMES, Inputs, assess  # noqa: E402
from waterhammer.model import simulate_moc  # noqa: E402

BOOK = ROOT / "excel" / "waterhammer.xlsx"
CASES = [
    dict(in_preset="copper_10x0.6", in_L=8, in_flow=6, in_valve="solenoid"),
    dict(in_preset="pex_12x2.0", in_L=8, in_flow=6, in_valve="single_lever_mixer"),
    dict(in_preset="mlcp_12x1.6", in_L=15, in_flow=4, in_valve="quarter_turn", in_T=60),
    dict(in_preset="copper_15x0.7", in_L=3, in_flow=4, in_valve="screw_down", in_p=4.0),
    dict(in_preset="custom", in_od=12, in_wall=2.0, in_mat="pex", in_al=0, in_L=10, in_flow=4,
         in_valve="custom", in_tc=0.0, in_restraint="expansion_joints"),
]


def cell_of(wb, nm):
    dest = next(iter(wb.defined_names[nm].destinations))
    return wb[dest[0]][dest[1].replace("$", "")]


def recalc(case: dict, tmp: Path) -> dict:
    wb = load_workbook(BOOK)
    for k, v in case.items():
        cell_of(wb, k).value = v
    src = tmp / "in.xlsx"
    wb.save(src)
    out = tmp / "out"
    out.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=600)
    res = load_workbook(out / "in.xlsx", data_only=True)
    names = ["x_a", "x_v", "x_djk", "x_dcf", "x_dmoc", "x_ddes", "x_pmin", "x_required", "x_varr", "x_rows_ok"]
    return {n: cell_of(res, n).value for n in names}


def python(case: dict) -> dict:
    if case["in_preset"] == "custom":
        from waterhammer import MATERIALS, Pipe
        pipe = Pipe(case["in_od"], case["in_wall"], MATERIALS[case["in_mat"]], aluminium_mm=case["in_al"])
    else:
        pipe = PRESETS[case["in_preset"]]
    tc = case["in_tc"] if case["in_valve"] == "custom" else VALVE_CLOSURE_TIMES[case["in_valve"]]
    inp = Inputs(pipe, case["in_L"], case["in_flow"], case.get("in_p", 3.0), case.get("in_T", 55), tc,
                 case.get("in_restraint", "anchored"), simulate=False)
    r = assess(inp)
    m = simulate_moc(pipe, inp.length_m, inp.flow_l_min / 60000, inp.static_pressure_bar * 1e5,
                     inp.temperature_c, tc, inp.restraint, reaches=10)
    return dict(x_a=r.wave_speed_m_s, x_v=r.velocity_m_s, x_djk=r.joukowsky_bar * 1e5,
                x_dcf=r.closed_form_surge_bar * 1e5, x_dmoc=m.peak_gauge_pa - inp.static_pressure_bar * 1e5,
                x_pmin=m.min_gauge_pa)


def main() -> int:
    if not shutil.which("soffice"):
        print("LibreOffice (soffice) not found")
        return 2
    worst = 0.0
    with tempfile.TemporaryDirectory() as d:
        for case in CASES:
            xl, py = recalc(case, Path(d)), python(case)
            print(case)
            for k, pv in py.items():
                xv = xl[k]
                rel = abs(xv - pv) / max(abs(pv), 1.0)
                worst = max(worst, rel)
                print(f"  {k:8} excel {xv:14.4f}  python {pv:14.4f}  rel diff {rel:.2e}")
            print(f"  verdict: {'MITIGATION REQUIRED' if xl['x_required'] else 'none'}; rows ok: {xl['x_rows_ok']}")
    print(f"worst relative difference {worst:.2e}")
    return 0 if worst < 1e-6 else 1


if __name__ == "__main__":
    raise SystemExit(main())
