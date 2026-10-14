"""Build deadleg.xlsx: the dead-leg model as live spreadsheet formulas.

    python excel/build_workbook.py [out.xlsx]

The workbook solves the same equations with the same numerics as
deadleg/model.py, on a fixed grid of 40 axial cells, 5 wall shells and 360
time steps:

* each step the water advances exactly one cell (dt = dx/u);
* radial exchange water -> 5 wall shells -> air is implicit; the tridiagonal
  system in each cell is solved by the Thomas algorithm laid out across
  sheets (m_1..m_5 forward sweep, dp_1..dp_5 right-hand side, x_0..x_5 back
  substitution), and the air node is eliminated exactly through two
  SUMPRODUCTs on the Air sheet;
* water-side and air-side heat transfer coefficients vary with the local
  temperatures (sheets hi and ho).

Axial conduction along the wall is left out; in the Python model it changes
delivery times by less than 0.01 s for these pipes.

Every sheet of the time-stepping grid has one row per time step (row 2 is
step 0) and one column per axial cell (B = cell 1 at the inlet, AO = cell 40
at the tap).
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NX = 40
NR = 5
STEPS = 360
FIRST = 2  # first data row
LAST = FIRST + STEPS - 1
COLS = [get_column_letter(i + 2) for i in range(NX)]  # B..AO
C1, CN = COLS[0], COLS[-1]

SERIF = "Palatino Linotype"
SANS = "Calibri"
INK = "222222"
GREY = "8C8C8C"
TUFTE_BG = "FFFFF8"
ECON_RED = "E3120B"
ECON_BLUE = "006BA2"
ECON_GRID = "B7C6CF"

MATERIALS = [
    # name, k, rho, cp, emissivity
    ("Copper", 340.0, 8940.0, 385.0, 0.1),
    ("PE-X", 0.38, 940.0, 2100.0, 0.9),
    ("MLCP", 0.45, 1150.0, 1750.0, 0.9),
]

wb = Workbook()


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


# ---------------------------------------------------------------- Inputs ----
inp = wb.active
inp.title = "Inputs"
inp.sheet_view.showGridLines = False
inp.column_dimensions["A"].width = 44
inp.column_dimensions["B"].width = 14
inp.column_dimensions["C"].width = 12
inp.column_dimensions["D"].width = 60

title_font = Font(name=SERIF, size=18, italic=True, color=INK)
head_font = Font(name=SERIF, size=12, italic=True, color=INK)
body_font = Font(name=SERIF, size=11, color=INK)
note_font = Font(name=SERIF, size=10, color="595959", italic=True)
input_fill = PatternFill("solid", fgColor="FFF7D6")
thin = Side(style="thin", color="BBBBBB")

inp["A1"] = "Hot water into a cold pipe: dead-leg model"
inp["A1"].font = title_font
inp["A2"] = ("Edit the shaded cells. Every other number in the workbook is a live formula. "
             "Results and charts are on the two Dashboard sheets.")
inp["A2"].font = note_font

rows = [
    ("Pipe", None, None, None),
    ("Pipe material", "Copper", "", "Copper, PE-X, MLCP or Custom (then fill in the custom properties below)"),
    ("Outside diameter", 10.0, "mm", ""),
    ("Wall thickness", 0.6, "mm", ""),
    ("Length of the run, L", 5.0, "m", ""),
    ("Operating conditions", None, None, None),
    ("Flow rate, Q", 3.0, "l/min", ""),
    ("Starting temperature of water, pipe and air, T0", 20.0, "°C", ""),
    ("Inlet water temperature, Tin", 55.0, "°C", "Constant from the moment the tap opens"),
    ("Air body", None, None, None),
    ("Air cross-section per metre of pipe", 0.01, "m²", "0.01 = a 100 × 100 mm void; 0.45 = a 1.5 × 0.3 m void"),
    ("Hold the air at T0 (a large room)?", "No", "", "Yes or No"),
    ("Heat loss from the air body, UA_env", 0.0, "W/K", "0 = enclosed and insulated"),
    ("Environment temperature, T_env", 20.0, "°C", "Only used when UA_env > 0"),
    ("Custom wall material (used only when material is Custom)", None, None, None),
    ("Conductivity, k_p", 0.4, "W/m K", ""),
    ("Density, ρ_p", 950.0, "kg/m³", ""),
    ("Specific heat, c_p", 2000.0, "J/kg K", ""),
    ("Outer-surface emissivity, ε", 0.9, "", ""),
    ("Thresholds for delivery time", None, None, None),
    ("Threshold 1", 40.0, "°C", ""),
    ("Threshold 2", 45.0, "°C", ""),
    ("Threshold 3", 50.0, "°C", ""),
]
names = {
    "Pipe material": "Material", "Outside diameter": "OD_mm", "Wall thickness": "Wall_mm",
    "Length of the run, L": "L", "Flow rate, Q": "Q_lpm",
    "Starting temperature of water, pipe and air, T0": "T_0", "Inlet water temperature, Tin": "Tin",
    "Air cross-section per metre of pipe": "AirArea", "Hold the air at T0 (a large room)?": "AirFixed",
    "Heat loss from the air body, UA_env": "UAenv", "Environment temperature, T_env": "Tenv",
    "Conductivity, k_p": "Custom_k", "Density, ρ_p": "Custom_rho", "Specific heat, c_p": "Custom_cp",
    "Outer-surface emissivity, ε": "Custom_eps", "Threshold 1": "Thr_1", "Threshold 2": "Thr_2", "Threshold 3": "Thr_3",
}
r = 4
for label, val, unit, note in rows:
    if val is None and unit is None:
        r += 1
        inp.cell(r, 1, label).font = head_font
        r += 1
        continue
    inp.cell(r, 1, label).font = body_font
    c = inp.cell(r, 2, val)
    c.font = Font(name=SERIF, size=11, bold=True, color=INK)
    c.fill = input_fill
    c.border = Border(bottom=thin)
    c.alignment = Alignment(horizontal="right")
    inp.cell(r, 3, unit).font = body_font
    inp.cell(r, 4, note).font = note_font
    name(names[label], f"Inputs!$B${r}")
    r += 1

dv = DataValidation(type="list", formula1='"Copper,PE-X,MLCP,Custom"', allow_blank=False)
inp.add_data_validation(dv)
dv.add(wb.defined_names["Material"].attr_text.split("!")[1].replace("$", ""))
dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
inp.add_data_validation(dv2)
dv2.add(wb.defined_names["AirFixed"].attr_text.split("!")[1].replace("$", ""))

# Material table
r += 1
inp.cell(r, 1, "Material properties (reference)").font = head_font
r += 1
hdr = ["Material", "k_p (W/m K)", "ρ_p (kg/m³)", "c_p (J/kg K)", "ε"]
for j, h in enumerate(hdr):
    inp.cell(r, 1 + j if j == 0 else 1 + j, h).font = Font(name=SERIF, size=10, italic=True, color=INK)
mat_top = r + 1
for m in MATERIALS:
    r += 1
    for j, v in enumerate(m):
        inp.cell(r, 1 + j, v).font = body_font
mat_bot = r
name("MatTable", f"Inputs!$A${mat_top}:$E${mat_bot}")
inp.cell(r + 2, 1, ("The model: plug flow of water; implicit heat exchange water → 5 wall shells → one well-mixed "
                    "air body; temperature-dependent film coefficients (Gnielinski / laminar blend inside, "
                    "Churchill–Chu natural convection plus radiation outside). 40 cells along the pipe, 360 time "
                    "steps of dt = dx/u. See the Python package deadleg and docs/physics/physics.pdf.")).font = note_font
inp.cell(r + 2, 1).alignment = Alignment(wrap_text=True, vertical="top")
inp.merge_cells(start_row=r + 2, start_column=1, end_row=r + 5, end_column=4)

# --------------------------------------------------------------- Derived ----
der = wb.create_sheet("Derived")
der.sheet_view.showGridLines = False
der.column_dimensions["A"].width = 46
der.column_dimensions["B"].width = 16
der.column_dimensions["C"].width = 12
der["A1"] = "Derived quantities (all formulas)"
der["A1"].font = title_font


def lookup(col):
    return f'IF(Material="Custom",{{c}},INDEX(MatTable,MATCH(Material,INDEX(MatTable,0,1),0),{col}))'


derived = [
    ("k_p", "Wall conductivity", lookup(2).format(c="Custom_k"), "W/m K"),
    ("rho_p", "Wall density", lookup(3).format(c="Custom_rho"), "kg/m³"),
    ("cp_p", "Wall specific heat", lookup(4).format(c="Custom_cp"), "J/kg K"),
    ("eps", "Outer-surface emissivity", lookup(5).format(c="Custom_eps"), ""),
    ("ro", "Outer radius r_o", "OD_mm/2000", "m"),
    ("ri", "Inner radius r_i", "ro-Wall_mm/1000", "m"),
    ("Di", "Bore diameter D_i", "2*ri", "m"),
    ("Do", "Outside diameter D_o", "2*ro", "m"),
    ("Ai", "Bore area A_i", "PI()*ri^2", "m²"),
    ("Ap", "Wall area A_p", "PI()*(ro^2-ri^2)", "m²"),
    ("dx", "Cell length dx = L/40", f"L/{NX}", "m"),
    ("u", "Mean water velocity u", "Q_lpm/60000/Ai", "m/s"),
    ("dt", "Time step dt = dx/u", "dx/u", "s"),
    ("Tref", "Temperature for water ρc", "(T_0+Tin)/2", "°C"),
    ("rhoCw", "Water ρ_w c_w at Tref",
     "1000*(1-(Tref+288.9414)/(508929.2*(Tref+68.12963))*(Tref-3.9863)^2)"
     "*(4217.4-3.720283*Tref+0.1412855*Tref^2-0.002654387*Tref^3+0.00002093236*Tref^4)", "J/m³ K"),
    ("Cw", "Water heat capacity per metre", "rhoCw*Ai", "J/m K"),
    ("Cwdt", "Cw/dt", "Cw/dt", "W/m K"),
    ("rhoCp", "Wall ρ_p c_p", "rho_p*cp_p", "J/m³ K"),
    ("sigma", "Heat capacity ratio σ", "rhoCp*Ap/Cw", ""),
    ("tplug", "Plug-flow time V/Q", "Ai*L/(Q_lpm/60000)", "s"),
    ("tfront", "Front time (1+σ)·t_plug", "(1+sigma)*tplug", "s"),
    ("Vwater", "Water in the pipe", "Ai*L*1000", "l"),
    ("Ca", "Air heat capacity", "IF(AirFixed=\"Yes\",0,101325/(287.05*(T_0+273.15))*1006*AirArea*L)", "J/K"),
    ("Cadt", "Ca/dt", "Ca/dt", "W/K"),
    ("Tmax", "Time covered by the 360 steps", f"{STEPS}*dt", "s"),
]
rr = 3
for nm, label, f, unit in derived:
    der.cell(rr, 1, label).font = body_font
    der.cell(rr, 2, "=" + f).font = body_font
    der.cell(rr, 2).number_format = "0.000E+00" if nm in ("ri", "ro", "Ai", "Ap", "Di", "Do") else "0.0000"
    der.cell(rr, 3, unit).font = body_font
    name(nm, f"Derived!$B${rr}")
    rr += 1

rr += 1
der.cell(rr, 1, "Wall shells (5 of equal thickness)").font = head_font
rr += 1
for j in range(NR + 1):
    der.cell(rr, 1, f"Face radius r_f{j}").font = body_font
    der.cell(rr, 2, f"=ri+(ro-ri)*{j}/{NR}")
    name(f"rf_{j}", f"Derived!$B${rr}")
    rr += 1
for j in range(1, NR + 1):
    der.cell(rr, 1, f"Shell {j} mid radius r_c{j}").font = body_font
    der.cell(rr, 2, f"=(rf_{j-1}+rf_{j})/2")
    name(f"rc_{j}", f"Derived!$B${rr}")
    rr += 1
for j in range(1, NR + 1):
    der.cell(rr, 1, f"Shell {j} heat capacity / dt, Cp{j}/dt").font = body_font
    der.cell(rr, 2, f"=rhoCp*PI()*(rf_{j}^2-rf_{j-1}^2)/dt")
    name(f"Cp{j}dt", f"Derived!$B${rr}")
    rr += 1
der.cell(rr, 1, "Conduction resistance, bore to shell 1 centre, R_in").font = body_font
der.cell(rr, 2, "=LN(rc_1/ri)/(2*PI()*k_p)")
name("Rin", f"Derived!$B${rr}")
rr += 1
der.cell(rr, 1, "Conduction resistance, shell 5 centre to outside, R_out").font = body_font
der.cell(rr, 2, f"=LN(ro/rc_{NR})/(2*PI()*k_p)")
name("Rout", f"Derived!$B${rr}")
rr += 1
for j in range(1, NR):
    der.cell(rr, 1, f"Conductance shell {j} → {j+1}, G_{j}{j+1}").font = body_font
    der.cell(rr, 2, f"=2*PI()*k_p/LN(rc_{j+1}/rc_{j})")
    name(f"G_{j}{j+1}", f"Derived!$B${rr}")
    rr += 1
for j in range(2, NR):
    der.cell(rr, 1, f"Constant diagonal, shell {j}, b{j}").font = body_font
    der.cell(rr, 2, f"=Cp{j}dt+G_{j-1}{j}+G_{j}{j+1}")
    name(f"bd_{j}", f"Derived!$B${rr}")
    rr += 1

# -------------------------------------------------------- formula sheets ----
WATER_RHO = "1000*(1-({T}+288.9414)/(508929.2*({T}+68.12963))*({T}-3.9863)^2)"
WATER_MU = "2.414E-5*10^(247.8/({T}+133.15))"
WATER_K = "(0.5706+0.001756*{T}-0.00000646*{T}^2)"
WATER_CP = "(4217.4-3.720283*{T}+0.1412855*{T}^2-0.002654387*{T}^3+0.00002093236*{T}^4)"


def hi_formula(T):
    re = f"(u*Di*{WATER_RHO.format(T=T)}/({WATER_MU.format(T=T)}))"
    pr = f"({WATER_MU.format(T=T)}*{WATER_CP.format(T=T)}/{WATER_K.format(T=T)})"

    def gn(R):
        f = f"(0.79*LN({R})-1.64)^-2"
        return f"(({f})/8*({R}-1000)*{pr}/(1+12.7*SQRT(({f})/8)*({pr}^(2/3)-1)))"

    nu = (f"IF({re}>=10000,{gn(re)},IF({re}<=2300,3.66,"
          f"(1-({re}-2300)/7700)*3.66+({re}-2300)/7700*{gn('10000')}))")
    return f"={nu}*{WATER_K.format(T=T)}/Di"


def ho_formula(Ts, Ta):
    tf = f"(({Ts})+({Ta}))/2"
    nu = f"(1.327E-5+9.2E-8*{tf})"
    k = f"(0.02414+0.000077*{tf})"
    ra = f"(9.81/({tf}+273.15)*ABS(({Ts})-({Ta}))*Do^3/({nu}*{nu}/0.71))"
    nud = f"(0.6+0.387*{ra}^(1/6)/(1+(0.559/0.71)^(9/16))^(8/27))^2"
    rad = f"eps*5.670374419E-8*((({Ts})+273.15)^2+(({Ta})+273.15)^2)*(({Ts})+({Ta})+546.3)"
    return f"={nud}*{k}/Do+{rad}"


def grid_sheet(title, note):
    ws = wb.create_sheet(title)
    ws["A1"] = note
    ws["A1"].font = Font(name=SERIF, size=9, italic=True, color="595959")
    for i, col in enumerate(COLS):
        ws[f"{col}1"] = i + 1
        ws[f"{col}1"].font = Font(name=SERIF, size=9, color=GREY)
    for k in range(STEPS):
        ws.cell(FIRST + k, 1, k).font = Font(name=SERIF, size=9, color=GREY)
    ws.freeze_panes = "B2"
    return ws


def prev(sheet, col, row):
    """End-of-previous-step value of a state sheet, or T0 on the first step."""
    return "T_0" if row == FIRST else f"{sheet}!{col}{row - 1}"


sheets = {}
sheets["Tw_in"] = grid_sheet("Tw_in", "Water temperature at the start of each step (after advection), °C. Row = step, column = cell.")
sheets["hi"] = grid_sheet("hi", "Water-side heat transfer coefficient h_i, W/m²K (Gnielinski / laminar blend at local T_w).")
sheets["ho"] = grid_sheet("ho", "Air-side coefficient h_o, W/m²K (Churchill–Chu natural convection + radiation).")
sheets["Gi"] = grid_sheet("Gi", "Conductance water → shell 1 centre per metre, W/m K.")
sheets["Go"] = grid_sheet("Go", "Conductance shell 5 centre → air per metre, W/m K.")
for j in range(1, NR + 1):
    sheets[f"m_{j}"] = grid_sheet(f"m_{j}", f"Thomas forward sweep: modified diagonal for shell {j}.")
for j in range(1, NR + 1):
    sheets[f"dp_{j}"] = grid_sheet(f"dp_{j}", f"Thomas forward sweep: modified right-hand side for shell {j}, °C.")
air = wb.create_sheet("Air")
for j in range(NR, 0, -1):
    sheets[f"x_{j}"] = grid_sheet(f"x_{j}", f"Wall shell {j} temperature at the end of each step, °C (shell 1 at the bore).")
sheets["x_0"] = grid_sheet("x_0", "Water temperature at the end of each step, before advection, °C. Column AO is the tap.")

for k in range(STEPS):
    row = FIRST + k
    for i, col in enumerate(COLS):
        # water entering the step: shifted from the previous step's end state
        if row == FIRST:
            twin = "=T_0"
        elif i == 0:
            twin = "=Tin"
        else:
            twin = f"=x_0!{COLS[i - 1]}{row - 1}"
        sheets["Tw_in"][f"{col}{row}"] = twin
        T = f"Tw_in!{col}{row}"
        sheets["hi"][f"{col}{row}"] = hi_formula(T)
        Ts = prev("x_5", col, row)
        Ta = f"Air!$B${row}"
        sheets["ho"][f"{col}{row}"] = ho_formula(Ts, Ta)
        sheets["Gi"][f"{col}{row}"] = f"=2*PI()*ri*hi!{col}{row}/(1+2*PI()*ri*hi!{col}{row}*Rin)"
        sheets["Go"][f"{col}{row}"] = f"=2*PI()*ro*ho!{col}{row}/(1+2*PI()*ro*ho!{col}{row}*Rout)"
        Gi, Go = f"Gi!{col}{row}", f"Go!{col}{row}"
        m = {j: f"m_{j}!{col}{row}" for j in range(1, NR + 1)}
        dp = {j: f"dp_{j}!{col}{row}" for j in range(1, NR + 1)}
        sheets["m_1"][f"{col}{row}"] = f"=Cp1dt+{Gi}+G_12-{Gi}^2/(Cwdt+{Gi})"
        for j in range(2, NR):
            sheets[f"m_{j}"][f"{col}{row}"] = f"=bd_{j}-G_{j-1}{j}^2/{m[j-1]}"
        sheets[f"m_{NR}"][f"{col}{row}"] = f"=Cp{NR}dt+G_{NR-1}{NR}+{Go}-G_{NR-1}{NR}^2/{m[NR-1]}"
        sheets["dp_1"][f"{col}{row}"] = f"=(Cp1dt*{prev('x_1', col, row)}+{Gi}*Cwdt*{T}/(Cwdt+{Gi}))/{m[1]}"
        for j in range(2, NR + 1):
            sheets[f"dp_{j}"][f"{col}{row}"] = f"=(Cp{j}dt*{prev(f'x_{j}', col, row)}+G_{j-1}{j}*{dp[j-1]})/{m[j]}"
        x = {j: f"x_{j}!{col}{row}" for j in range(NR + 1)}
        sheets[f"x_{NR}"][f"{col}{row}"] = f"={dp[NR]}+{Go}/{m[NR]}*Air!$E${row}"
        for j in range(NR - 1, 0, -1):
            sheets[f"x_{j}"][f"{col}{row}"] = f"={dp[j]}+G_{j}{j+1}/{m[j]}*{x[j+1]}"
        sheets["x_0"][f"{col}{row}"] = f"=(Cwdt*{T}+{Gi}*{x[1]})/(Cwdt+{Gi})"

# Air: implicit air temperature from two sums over the cells
air["A1"] = "Step"
for j, h in enumerate(["Ta at start, °C", "Σ Go·X0 (W/m K·°C)", "Σ Go·(1−X1) (W/m K)", "Ta at end, °C"]):
    air.cell(1, 2 + j, h)
for c in range(1, 6):
    air.cell(1, c).font = Font(name=SERIF, size=9, italic=True, color="595959")
    air.column_dimensions[get_column_letter(c)].width = 22
for k in range(STEPS):
    row = FIRST + k
    air.cell(row, 1, k)
    air.cell(row, 2, "=T_0" if row == FIRST else f"=E{row - 1}")
    rng = lambda s: f"{s}!{C1}{row}:{CN}{row}"
    air.cell(row, 3, f"=SUMPRODUCT({rng('Go')},{rng(f'dp_{NR}')})")
    air.cell(row, 4, f"=SUMPRODUCT({rng('Go')},1-{rng('Go')}/{rng(f'm_{NR}')})")
    air.cell(row, 5, f'=IF(AirFixed="Yes",B{row},(Cadt*B{row}+dx*C{row}+UAenv*Tenv)/(Cadt+dx*D{row}+UAenv))')
air.freeze_panes = "B2"

# ---------------------------------------------------------------- Results ----
res = wb.create_sheet("Results", 2)
res.sheet_view.showGridLines = False
heads = ["Step", "Time, s", "Inlet, °C", "Outlet, °C", "Air, °C", "Water run off, l",
         "Crossing 1, s", "Crossing 2, s", "Crossing 3, s"]
for j, h in enumerate(heads):
    c = res.cell(1, 1 + j, h)
    c.font = Font(name=SERIF, size=10, italic=True, color=INK)
    res.column_dimensions[get_column_letter(1 + j)].width = 14
for k in range(STEPS):
    row = FIRST + k
    res.cell(row, 1, k)
    res.cell(row, 2, f"=({k}+0.5)*dt")
    res.cell(row, 3, "=Tin")
    res.cell(row, 4, f"=x_0!{CN}{row}")
    res.cell(row, 5, f"=Air!E{row}")
    res.cell(row, 6, f"=({k}+1)*Ai*dx*1000")
    for j, thr in enumerate(("Thr_1", "Thr_2", "Thr_3")):
        if row == FIRST:
            f = f'=IF(D{row}>={thr},B{row},"")'
        else:
            f = f'=IF(AND(D{row}>={thr},D{row - 1}<{thr}),B{row - 1}+({thr}-D{row - 1})*(B{row}-B{row - 1})/(D{row}-D{row - 1}),"")'
        res.cell(row, 7 + j, f)
    for j in range(1, 10):
        res.cell(row, j).font = Font(name=SERIF, size=10, color=INK)
        res.cell(row, j).number_format = "0.000" if j > 1 else "0"
res.freeze_panes = "A2"
name("ResT", f"Results!$B${FIRST}:$B${LAST}")
name("ResOut", f"Results!$D${FIRST}:$D${LAST}")


# ------------------------------------------------------------- Dashboards ----
def metric_rows():
    return [
        ("Pipe", '=Material&" "&TEXT(OD_mm,"0.0")&" × "&TEXT(Wall_mm,"0.0#")&" mm, "&TEXT(L,"0.#")&" m"', "", "@"),
        ("Flow", "=Q_lpm", "l/min", "0.0"),
        ("Water held in the pipe", "=Vwater", "l", "0.000"),
        ("Plug-flow time, V/Q", "=tplug", "s", "0.0"),
        ("Heat capacity ratio, σ", "=sigma", "", "0.00"),
        ("Front time, (1+σ)·V/Q", "=tfront", "s", "0.0"),
        ('="Time to "&TEXT(Thr_1,"0")&" °C"', f'=IFERROR(MIN(Results!G{FIRST}:G{LAST})/(COUNT(Results!G{FIRST}:G{LAST})>0),"not reached")', "s", "0.0"),
        ('="Time to "&TEXT(Thr_2,"0")&" °C"', f'=IFERROR(MIN(Results!H{FIRST}:H{LAST})/(COUNT(Results!H{FIRST}:H{LAST})>0),"not reached")', "s", "0.0"),
        ('="Time to "&TEXT(Thr_3,"0")&" °C"', f'=IFERROR(MIN(Results!I{FIRST}:I{LAST})/(COUNT(Results!I{FIRST}:I{LAST})>0),"not reached")', "s", "0.0"),
        ('="Water run off before "&TEXT(Thr_3,"0")&" °C"', "=IFERROR(D14*Q_lpm/60,\"not reached\")", "l", "0.00"),
        ("Outlet at end of run", f"=Results!D{LAST}", "°C", "0.0"),
        ("Air at end of run", f"=Results!E{LAST}", "°C", "0.0"),
        ("Time covered", "=Tmax", "s", "0.0"),
    ]


def chart_text(font_face, size=900, color=INK, bold=False):
    cp = CharacterProperties(latin=DFont(typeface=font_face), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def line(color, width_pt, dash=None):
    ln = LineProperties(solidFill=color, w=int(width_pt * 12700))
    if dash:
        ln.prstDash = dash
    return GraphicalProperties(ln=ln)


def build_dashboard(title, theme):
    ws = wb.create_sheet(title, 1)
    ws.sheet_view.showGridLines = False
    tufte = theme == "tufte"
    face = SERIF if tufte else SANS
    bg = PatternFill("solid", fgColor=TUFTE_BG if tufte else "FFFFFF")
    for rr_ in range(1, 45):
        for cc in range(1, 16):
            ws.cell(rr_, cc).fill = bg
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 8
    if not tufte:
        red = PatternFill("solid", fgColor=ECON_RED)
        for cc in range(1, 16):
            ws.cell(1, cc).fill = red
        ws.row_dimensions[1].height = 4
        ws["B2"].fill = red
        ws.row_dimensions[2].height = 10
    ws["B3"] = "Hot water into a cold pipe" if tufte else "Waiting for hot water"
    ws["B3"].font = Font(name=face, size=20, italic=tufte, bold=not tufte, color=INK)
    ws["B4"] = ('="Outlet temperature after the tap opens: "&Material&" pipe, "&TEXT(L,"0.#")&" m at "'
                '&TEXT(Q_lpm,"0.#")&" l/min, inlet "&TEXT(Tin,"0")&" °C, pipe initially "&TEXT(T_0,"0")&" °C"')
    ws["B4"].font = Font(name=face, size=11, italic=tufte, color=INK)
    r0 = 6
    for j, (label, f, unit, fmt) in enumerate(metric_rows()):
        rr_ = r0 + j
        a = ws.cell(rr_, 2, label)
        a.font = Font(name=face, size=11, color=INK)
        ws.merge_cells(start_row=rr_, start_column=2, end_row=rr_, end_column=3)
        v = ws.cell(rr_, 4, f)
        v.font = Font(name=face, size=11, bold=not tufte, color=INK)
        v.number_format = fmt
        v.alignment = Alignment(horizontal="left" if fmt == "@" else "right")
        ws.cell(rr_, 5, unit).font = Font(name=face, size=10, color="595959")
        if not tufte:
            for cc in (2, 3, 4, 5):
                ws.cell(rr_, cc).border = Border(bottom=Side(style="thin", color=ECON_GRID))
    ws.cell(r0 + 15, 2, ("Theme: Tufte. For the Economist style, see the other dashboard sheet."
                         if tufte else "Theme: The Economist. For the Tufte style, see the other dashboard sheet.")
            ).font = Font(name=face, size=9, italic=True, color="595959")
    ws.cell(r0 + 16, 2, "Source: deadleg model, 40 cells × 5 wall shells, live spreadsheet formulas").font = Font(
        name=face, size=9, italic=tufte, color="595959")

    ch = ScatterChart()
    ch.style = 13
    ch.height = 10.5
    ch.width = 19
    ch.legend = None if tufte else ch.legend
    if not tufte:
        ch.legend.position = "t"
    x = Reference(res, min_col=2, min_row=FIRST, max_row=LAST)
    for col, label, color, w, dash in (
        (4, "Outlet", INK if tufte else ECON_BLUE, 1.75, None),
        (3, "Inlet", GREY if tufte else "3EBCD2", 0.75, "dash"),
        (5, "Air", GREY if tufte else "758D99", 0.75, None),
    ):
        y = Reference(res, min_col=col, min_row=FIRST, max_row=LAST)
        s = Series(y, x, title=label)
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties = line(color, w, dash)
        ch.series.append(s)
    ch.x_axis.title = "Seconds after the tap opens"
    ch.y_axis.title = "°C"
    ch.x_axis.number_format = "0"
    ch.y_axis.number_format = "0"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.majorTickMark = "out"
    ch.y_axis.majorTickMark = "out" if tufte else "none"
    ch.x_axis.txPr = chart_text(face)
    ch.y_axis.txPr = chart_text(face)
    if tufte:
        ch.y_axis.majorGridlines = None
        ch.x_axis.majorGridlines = None
        ch.x_axis.graphicalProperties = line(INK, 0.5)
        ch.y_axis.graphicalProperties = line(INK, 0.5)
    else:
        ch.x_axis.majorGridlines = None
        ch.y_axis.majorGridlines.spPr = line(ECON_GRID, 0.75)
        ch.y_axis.crosses = "max" if False else ch.y_axis.crosses
        ch.x_axis.graphicalProperties = line(INK, 0.75)
        ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ws.add_chart(ch, "G3")
    if tufte:
        ws.cell(26, 7, "Outlet (black), inlet (dashed grey) and air (thin grey) temperatures.").font = Font(
            name=face, size=9, italic=True, color="595959")
    return ws


build_dashboard("Dashboard (Economist)", "economist")
build_dashboard("Dashboard (Tufte)", "tufte")

# order: Dashboard (Tufte), Dashboard (Economist), Inputs, Results, Derived, then grids
order = ["Dashboard (Tufte)", "Dashboard (Economist)", "Inputs", "Results", "Derived"]
wb._sheets = [wb[n] for n in order] + [s for s in wb._sheets if s.title not in order]
wb.active = 0

out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("deadleg.xlsx")
wb.save(out)
print("wrote", out)
