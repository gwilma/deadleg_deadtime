"""Cross-check the Excel workbook against the Python model (and the JS port).

For each case the inputs of excel/deadleg.xlsx are set with openpyxl, the
workbook is recalculated by LibreOffice (headless), and the outlet and air
temperatures it computes are compared step by step with Python and JS run on
the workbook's grid: 40 cells, 5 wall shells, no axial wall conduction, 360
steps of dt = dx/u.

    python crosscheck/compare_excel.py [excel/deadleg.xlsx] [--out crosscheck/excel_vs_python.json]

Needs LibreOffice (soffice) with Calc, and node for the JS column.
"""
from __future__ import annotations

import argparse
import json
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import openpyxl

from deadleg import PRESETS, Scenario, make_pipe, simulate
from deadleg.properties import COPPER, MLCP, PEX

HERE = Path(__file__).resolve().parent
NX, NR, STEPS = 40, 5, 360
MAT = {"Copper": COPPER, "PE-X": PEX, "MLCP": MLCP}
PRESET_KEY = {("Copper", 10, 0.6): "copper_10x0.6", ("Copper", 12, 0.6): "copper_12x0.6",
              ("PE-X", 10, 1.5): "pex_10x1.5", ("PE-X", 12, 2.0): "pex_12x2.0",
              ("MLCP", 12, 1.6): "mlcp_12x1.6"}

CASES = [
    dict(id="copper10_3lpm", Material="Copper", OD_mm=10, Wall_mm=0.6, L=5, Q_lpm=3, T_0=20, Tin=55,
         AirArea=0.01, AirFixed="No", UAenv=0, Tenv=20),
    dict(id="pex12_1lpm", Material="PE-X", OD_mm=12, Wall_mm=2.0, L=5, Q_lpm=1, T_0=20, Tin=55,
         AirArea=0.01, AirFixed="No", UAenv=0, Tenv=20),
    dict(id="copper12_6lpm_room", Material="Copper", OD_mm=12, Wall_mm=0.6, L=8, Q_lpm=6, T_0=15, Tin=60,
         AirArea=0.01, AirFixed="Yes", UAenv=0, Tenv=15),
    dict(id="mlcp12_4lpm_loss", Material="MLCP", OD_mm=12, Wall_mm=1.6, L=15, Q_lpm=4, T_0=19, Tin=50,
         AirArea=0.45, AirFixed="No", UAenv=2, Tenv=15),
]


def set_inputs(src: Path, dst: Path, case: dict) -> None:
    wb = openpyxl.load_workbook(src)
    for k, v in case.items():
        if k == "id":
            continue
        sheet, ref = next(iter(wb.defined_names[k].destinations))
        wb[sheet][ref.replace("$", "")] = v
    wb.save(dst)


def recalc(path: Path, outdir: Path) -> Path:
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir), str(path)],
                   check=True, capture_output=True)
    return outdir / path.name


def read_excel(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    res = wb["Results"]
    t, out, air = [], [], []
    for row in res.iter_rows(min_row=2, max_row=1 + STEPS, min_col=2, max_col=5, values_only=True):
        t.append(row[0]); out.append(row[2]); air.append(row[3])
    der = {}
    for n in ("dt", "tplug", "sigma", "Vwater"):
        sheet, ref = next(iter(wb.defined_names[n].destinations))
        der[n] = wb[sheet][ref.replace("$", "")].value
    return dict(t=np.array(t, float), Tout=np.array(out, float), Tair=np.array(air, float), **der)


def pipe_for(case):
    key = PRESET_KEY.get((case["Material"], case["OD_mm"], case["Wall_mm"]))
    if key:
        return PRESETS[key], key
    return make_pipe(case["OD_mm"] / 1000, case["Wall_mm"] / 1000, MAT[case["Material"]]), None


def run_python(case, dt):
    pipe, _ = pipe_for(case)
    air = math.inf if case["AirFixed"] == "Yes" else case["AirArea"] * case["L"]
    sc = Scenario(pipe, length=case["L"], flow=case["Q_lpm"], duration=STEPS * dt,
                  T_initial=case["T_0"], T_inlet=case["Tin"], air_volume=air,
                  air_loss_UA=case["UAenv"], T_environment=case["Tenv"],
                  n_axial=NX, n_radial=NR, wall_axial_conduction=False)
    return simulate(sc)


def run_js(cases):
    js_cases = []
    for c in cases:
        _, key = pipe_for(c)
        js_cases.append(dict(id=c["id"], pipe=key or dict(od=c["OD_mm"] / 1000, wall=c["Wall_mm"] / 1000,
                                                          material={"Copper": "copper", "PE-X": "pex", "MLCP": "mlcp"}[c["Material"]]),
                             length=c["L"], flow=c["Q_lpm"], T0=c["T_0"], Tin=c["Tin"],
                             airArea="inf" if c["AirFixed"] == "Yes" else c["AirArea"],
                             airUA=c["UAenv"], Tenv=c["Tenv"], duration=1e9, maxSteps=STEPS))
    with tempfile.TemporaryDirectory() as d:
        cin, cout = Path(d, "cases.json"), Path(d, "out.json")
        cin.write_text(json.dumps(js_cases))
        subprocess.run(["node", str(HERE / "run_js.js"), str(cin), str(cout), str(NX), str(NR), "0"], check=True)
        return json.loads(cout.read_text())


def crossing(t, T, thr):
    i = np.flatnonzero(T >= thr)
    if not len(i):
        return math.nan
    i = i[0]
    if i == 0:
        return float(t[0])
    return float(t[i - 1] + (thr - T[i - 1]) * (t[i] - t[i - 1]) / (T[i] - T[i - 1]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook", nargs="?", default=str(HERE.parent / "excel" / "deadleg.xlsx"))
    ap.add_argument("--out", default=str(HERE / "excel_vs_python.json"))
    a = ap.parse_args()
    if not shutil.which("soffice"):
        sys.exit("soffice (LibreOffice Calc) is needed to recalculate the workbook")
    js = run_js(CASES)
    report = []
    with tempfile.TemporaryDirectory() as d:
        d = Path(d)
        for c in CASES:
            src = d / f"{c['id']}.xlsx"
            set_inputs(Path(a.workbook), src, c)
            (d / "calc").mkdir(exist_ok=True)
            xl = read_excel(recalc(src, d / "calc"))
            py = run_python(c, xl["dt"])
            n = min(len(py.T_out), STEPS)
            j = js[c["id"]]
            jt, jo, ja = (np.array(j[k][:n]) for k in ("t", "Tout", "Tair"))
            row = dict(id=c["id"], steps=n, dt=xl["dt"],
                       max_dTout_excel_python=float(np.max(np.abs(xl["Tout"][:n] - py.T_out[:n]))),
                       max_dTair_excel_python=float(np.max(np.abs(xl["Tair"][:n] - py.T_air[:n]))),
                       max_dTout_excel_js=float(np.max(np.abs(xl["Tout"][:n] - jo))),
                       max_dTout_js_python=float(np.max(np.abs(jo - py.T_out[:n]))),
                       max_dt_excel_python=float(np.max(np.abs(xl["t"][:n] - py.t[:n]))))
            for thr in (40, 45, 50):
                row[f"t{thr}"] = dict(excel=crossing(xl["t"], xl["Tout"], thr), python=py.time_to_reach(thr),
                                      js=crossing(jt, jo, thr))
            report.append(row)
            print(f"{c['id']:22s} max|dTout| Excel-Python {row['max_dTout_excel_python']:.2e} K, "
                  f"Excel-JS {row['max_dTout_excel_js']:.2e} K, max|dTair| {row['max_dTair_excel_python']:.2e} K; "
                  + ", ".join(f"t{thr} {row[f't{thr}']['excel']:.2f}/{row[f't{thr}']['python']:.2f}/{row[f't{thr}']['js']:.2f} s"
                              for thr in (40, 45, 50)))
    Path(a.out).write_text(json.dumps(report, indent=1))
    print("wrote", a.out)


if __name__ == "__main__":
    main()
